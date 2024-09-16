from sqlalchemy.orm import sessionmaker
from sqlalchemy import create_engine
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# Importar a configuração do banco de dados
from config import Config

# Configuração do banco de dados usando o arquivo config.py
engine = create_engine(Config.SQLALCHEMY_DATABASE_URI)
Session = sessionmaker(bind=engine)
session = Session()

# Importar os modelos
from models import db, Projeto, Tarefa, Superintendentes, Funcionarios, ProjetoSuperintendencia, ProjetoKeyUser, ProjetoFuncTI, Funcionario_TI


# Criação de uma lista para armazenar os dados
dados = []

# Consulta os dados
projetos = session.query(Projeto).all()

for projeto in projetos:
    # Consulta superintendentes e funcionários associados
    superintendencias = session.query(Superintendentes).join(ProjetoSuperintendencia).filter(ProjetoSuperintendencia.projeto_id == projeto.id).all()
    funcionarios = session.query(Funcionarios).join(ProjetoKeyUser).filter(ProjetoKeyUser.projeto_id == projeto.id).all()
    funcionarios_ti = session.query(Funcionarios).join(ProjetoFuncTI).filter(ProjetoFuncTI.projeto_id == projeto.id).all()

    for superintendente in superintendencias:
        for funcionario in funcionarios:
            for funcionario_ti in funcionarios_ti:
                dados.append({
                    'ID': projeto.id,
                    'Nome do Projeto': projeto.nome_projeto,
                    'Data Relatório': projeto.data_relatorio,
                    'Project Status': projeto.project_status,
                    'Status Geral': projeto.status_geral,
                    'ID Jira': projeto.id_jira,
                    'Order Number': projeto.order_number if projeto.order_number else 0,
                    'Nome Superintendente': superintendente.nomesuperintendente,
                    'Email Superintendente': superintendente.emailsuperintendente,
                    'Nome Funcionário': funcionario.nome,
                    'Email Funcionário': funcionario.email,
                    'Nome Funcionário TI': funcionario_ti.nome,  # Ajuste o nome do campo conforme seu modelo
                    'Email Funcionário TI': funcionario_ti.email  # Ajuste o nome do campo conforme seu modelo
                })
# Criação do DataFrame
df = pd.DataFrame(dados)

# Exportação para Excel
excel_path = 'relatorio_projetos.xlsx'
df.to_excel(excel_path, index=False, engine='openpyxl')

# Carregar o arquivo Excel e adicionar a tabela
wb = load_workbook(excel_path)
ws = wb.active

# Definir o intervalo da tabela
table = Table(displayName="TabelaProjetos", ref=ws.dimensions)

# Adicionar estilo à tabela
style = TableStyleInfo(
    name="TableStyleMedium9",
    showFirstColumn=False,
    showLastColumn=False,
    showRowStripes=True,
    showColumnStripes=True
)
table.tableStyleInfo = style

# Adicionar a tabela à planilha
ws.add_table(table)

# Salvar o arquivo Excel
wb.save(excel_path)

print("Relatório Excel gerado com sucesso!")